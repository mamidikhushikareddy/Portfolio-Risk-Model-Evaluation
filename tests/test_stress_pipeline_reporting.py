import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from riskplatform import stress as st
from riskplatform.config import TICKERS
from riskplatform.pipeline import run_analysis
from riskplatform.reporting import export_excel, render_brief
from riskplatform.validation import validate_prices


def _event_prices():
    idx = pd.bdate_range("2020-01-01", periods=10)
    return pd.DataFrame({"A": [100, 100, 90, 80, 70, 75, 85, 95, 100, 110.0],
                         "B": [100, 100, 100, 100, 100, 100, 100, 100, 100, 100.0]}, index=idx)


def test_replay_known_values_and_recovery():
    p = _event_prices()
    w = pd.Series({"A": 0.5, "B": 0.5})
    r = st.replay_event(p, w, "ev", str(p.index[2].date()), str(p.index[4].date()))
    assert r["portfolio_available"]
    # base = index[1] (100); end = index[4] (70): A -30%, B 0%
    assert r["asset_table"].loc["A", "event_return"] == pytest.approx(-0.30)
    assert r["portfolio_return_buy_and_hold"] == pytest.approx(-0.15)
    assert r["portfolio_return_daily_rebalanced"] == pytest.approx(
        (1 + 0.5 * (-0.1) )*(1 + 0.5 * (80/90 - 1))*(1 + 0.5 * (70/80 - 1)) - 1)
    assert r["trough_date"] == p.index[4]
    assert r["recovery_date"] is not None and r["days_trough_to_recovery"] > 0


def test_replay_withholds_when_data_missing():
    p = _event_prices()
    p.loc[p.index[:4], "B"] = np.nan  # B starts late
    r = st.replay_event(p, pd.Series({"A": 0.5, "B": 0.5}), "ev", str(p.index[2].date()), str(p.index[4].date()))
    assert not r["portfolio_available"] and "withheld" in r["reason"]
    assert r["asset_table"].loc["B", "status"].startswith("unavailable")
    assert r["asset_table"].loc["A", "event_return"] == pytest.approx(-0.30)  # observed asset still shown


def test_replay_outside_data_range():
    p = _event_prices()
    r = st.replay_event(p, pd.Series({"A": 1.0}), "old", "2005-01-03", "2005-03-01")
    assert not r["portfolio_available"] and r["reason"]
    r2 = st.replay_event(p, pd.Series({"A": 1.0}), "late", "2020-01-08", "2021-01-01")
    assert not r2["portfolio_available"]


def test_hypothetical_requires_explicit_shocks():
    w = pd.Series({"A": 0.6, "B": 0.4})
    h = st.hypothetical_shock(w, {"A": -0.2, "B": 0.1}, "x")
    assert h["portfolio_return"] == pytest.approx(-0.08) and "HYPOTHETICAL" in h["kind"]
    with pytest.raises(ValueError, match="no silent zero"):
        st.hypothetical_shock(w, {"A": -0.2})


def test_duration_approximation_known():
    w = pd.Series({"AGG": 0.5, "TLT": 0.2, "VTI": 0.3})
    d = st.duration_approximation(w, 100, {"AGG": 6.0, "TLT": 17.0}, {"AGG": 60.0, "TLT": 300.0})
    agg = -6 * 0.01 + 0.5 * 60 * 0.01 ** 2
    assert d["table"].loc["AGG", "approx_price_return"] == pytest.approx(agg)
    assert d["bond_contribution_to_portfolio"] == pytest.approx(0.5 * agg + 0.2 * (-0.17 + 0.5 * 300 * 1e-4))
    assert "unmodelled" in d["caveat"]


def test_yield_context_uses_levels_only():
    idx = pd.bdate_range("2020-01-01", periods=10)
    y = pd.DataFrame({"DGS10": [1.0, np.nan, 1.2, 1.3, np.nan, 1.5, 1.6, 1.7, 1.8, 1.9]}, index=idx)
    ctx = st.yield_context(y, str(idx[2].date()), str(idx[6].date()))
    assert ctx.loc["DGS10", "change_pp"] == pytest.approx(1.6 - 1.0)


def test_pipeline_end_to_end_and_weight_changes(prices, weights):
    res = run_analysis(prices, weights, window=500)
    assert len(res.backtest_summary) == 4
    assert res.risk_contrib["pct_contribution"].sum() == pytest.approx(1.0)
    assert res.portfolio_stats["ann_volatility"] == pytest.approx(res.portfolio_stats["ann_volatility_cov_check"])
    assert all(r["portfolio_available"] for r in res.stress_replays)
    bonds = pd.Series({t: 0.0 for t in TICKERS}); bonds["AGG"] = 1.0
    res2 = run_analysis(prices, bonds, run_backtests=False)
    assert res2.portfolio_stats["ann_volatility"] < res.portfolio_stats["ann_volatility"]
    assert res2.risk_contrib.loc["AGG", "pct_contribution"] == pytest.approx(1.0)
    assert res2.risk_contrib.loc["VTI", "weight"] == 0.0


def test_pipeline_rejects_bad_weights_and_missing_assets(prices, weights):
    bad = weights.copy(); bad["VTI"] += 0.05
    with pytest.raises(ValueError, match="sum to"):
        run_analysis(prices, bad)
    with pytest.raises(KeyError):
        run_analysis(prices.drop(columns=["GLD"]), weights)
    with pytest.raises(ValueError):
        run_analysis(prices, pd.Series({t: 0.0 for t in TICKERS}))


def test_short_sample_skips_backtest(prices, weights):
    res = run_analysis(prices.iloc[:300], weights, window=500)
    assert res.backtest_summary.empty


def test_late_listing_truncates_sample_without_filling(prices, weights):
    p = prices.copy()
    p.loc[p.index[:800], "EEM"] = np.nan
    res = run_analysis(p, weights, run_backtests=False)
    assert res.sample_info["start"] >= p.index[800]
    assert res.sample_info["rows_dropped_missing"] >= 800


def test_excel_export_deterministic(tmp_path, prices, weights, yields):
    q = validate_prices(prices, list(prices.columns))
    res = run_analysis(prices, weights, yields=yields, quality=q,
                       metadata={"source": "SYNTHETIC", "synthetic": True, "retrieved_at_utc": "n/a"})
    out = export_excel(res, tmp_path / "r.xlsx")
    wb = load_workbook(out)
    expected = {"Executive Summary", "Weights", "Asset Statistics", "Correlation Matrix", "Risk Contributions",
                "VaR and ES", "Backtesting", "Stress Scenarios", "Data Quality", "Methodology"}
    assert expected == set(wb.sheetnames)
    assert "SYNTHETIC" in wb["Executive Summary"]["A3"].value
    ws = wb["Weights"]
    assert any(isinstance(c.value, str) and c.value.startswith("=SUM") for row in ws.iter_rows() for c in row)
    es = wb["Executive Summary"]
    vol_cell = next(r for r in es.iter_rows(min_row=5) if r[0].value == "Annualised volatility")[1]
    assert vol_cell.value == pytest.approx(res.portfolio_stats["ann_volatility"]) and vol_cell.number_format == "0.00%"
    rc = wb["Risk Contributions"]
    pct = [r[4].value for r in rc.iter_rows(min_row=6, max_row=5 + len(TICKERS))]
    assert sum(pct) == pytest.approx(1.0)
    assert wb["Correlation Matrix"].conditional_formatting


def test_brief_refuses_synthetic_and_populates_real(prices, weights, tmp_path):
    tpl = tmp_path / "t.md"
    tpl.write_text("Vol {{ANN_VOL}} from {{SAMPLE_START}} top {{TOP_RISK_ASSET}}\n{{BACKTEST_TABLE}}")
    res = run_analysis(prices, weights, metadata={"synthetic": True})
    with pytest.raises(ValueError, match="SYNTHETIC"):
        render_brief(res, tpl)
    res.metadata = {"synthetic": False, "source": "test"}  # stands in for verified real data
    text = render_brief(res, tpl)
    assert "{{" not in text and "%" in text
