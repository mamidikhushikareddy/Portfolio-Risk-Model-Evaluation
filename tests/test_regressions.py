"""Regression tests for defects found in review (each failed before it was fixed)."""
import json
import sys
from pathlib import Path

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

from riskplatform import analytics as an
from riskplatform import cli
from riskplatform import ingestion as ing
from riskplatform.pipeline import run_analysis
from riskplatform.reporting import export_excel
from riskplatform.validation import validate_prices, validate_yields

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "examples"))


# ---- cache must not silently go stale ----
def _series(n=3):
    return pd.Series(np.arange(n) + 1.0, index=pd.bdate_range("2020-01-01", periods=n))


def test_open_ended_cache_older_than_today_triggers_redownload(tmp_path):
    ing.download_prices(["X"], "2020-01-01", None, cache_dir=tmp_path, downloader=lambda t, s, e: _series())
    meta = tmp_path / "X.meta.json"
    j = json.loads(meta.read_text()); j["retrieved_at_utc"] = "2020-06-01T00:00:00+00:00"; meta.write_text(json.dumps(j))
    calls = []
    def dl(t, s, e):
        calls.append(t); return _series(4)
    r = ing.download_prices(["X"], "2020-01-01", None, cache_dir=tmp_path, downloader=dl)
    assert calls == ["X"] and r.metadata["per_ticker"]["X"]["origin"] == "download" and len(r.prices) == 4


def test_old_cache_with_failed_download_is_flagged_stale(tmp_path):
    ing.download_prices(["X"], "2020-01-01", None, cache_dir=tmp_path, downloader=lambda t, s, e: _series())
    meta = tmp_path / "X.meta.json"
    j = json.loads(meta.read_text()); j["retrieved_at_utc"] = "2020-06-01T00:00:00+00:00"; meta.write_text(json.dumps(j))
    def offline(t, s, e):
        raise RuntimeError("down")
    r = ing.download_prices(["X"], "2020-01-01", None, cache_dir=tmp_path, downloader=offline)
    assert r.metadata["per_ticker"]["X"]["origin"] == "stale_cache" and "warning" in r.metadata["per_ticker"]["X"]


def test_closed_historical_range_cache_is_reused(tmp_path):
    ing.download_prices(["X"], "2020-01-01", "2020-01-03", cache_dir=tmp_path, downloader=lambda t, s, e: _series())
    def boom(*a):
        raise AssertionError("should not download")
    r = ing.download_prices(["X"], "2020-01-01", "2020-01-03", cache_dir=tmp_path, downloader=boom)
    assert r.metadata["per_ticker"]["X"]["origin"] == "cache"


def test_yfinance_end_date_is_inclusive(monkeypatch):
    import yfinance as yf
    seen = {}
    def fake(ticker, start, end, **kw):
        seen["end"] = end
        return pd.DataFrame({"Close": [1.0, 2.0]}, index=pd.bdate_range("2020-01-02", periods=2))
    monkeypatch.setattr(yf, "download", fake)
    ing._yfinance_download("X", "2020-01-01", "2020-01-31")
    assert seen["end"] == "2020-02-01"  # yfinance `end` is exclusive


# ---- drawdown must not bridge data gaps ----
def test_asset_drawdown_uses_prices_across_gap():
    idx = pd.bdate_range("2020-01-01", periods=6)
    p = pd.Series([100, 100, np.nan, 50, 60, 70.0], index=idx)
    assert an.max_drawdown_from_prices(p) == pytest.approx(-0.5)
    # the returns-only path would miss the 100 -> 50 move that spans the gap
    assert an.drawdown_stats(an.simple_returns(p).dropna())["max_drawdown"] > -0.2


# ---- yields are validated too ----
def test_yield_validation():
    idx = pd.bdate_range("2020-01-01", periods=5)
    y = pd.DataFrame({"DGS10": [1.5, np.nan, 1.6, 40.0, 1.7]}, index=idx)
    v = validate_yields(y)
    assert v["series"]["DGS10"]["n_missing"] == 1
    assert any(i["code"] == "IMPLAUSIBLE_YIELD" for i in v["issues"])
    assert validate_yields(pd.DataFrame())["issues"][0]["code"] == "NO_YIELDS"
    rep = validate_prices(pd.DataFrame({"A": np.linspace(1, 2, 50)}, index=pd.bdate_range("2020-01-01", periods=50)), yields=y)
    assert "DGS10" in rep["yields"]


# ---- Excel: every numeric cell carries an explicit, correct format ----
def test_excel_number_formats(tmp_path, prices, weights):
    res = run_analysis(prices, weights, metadata={"synthetic": True, "source": "SYNTHETIC"})
    wb = load_workbook(export_excel(res, tmp_path / "f.xlsx"))
    ws = wb["Backtesting"]
    hdr = {c.value: c.column for c in ws[5]}
    first = 6
    assert ws.cell(first, hdr["mean_forecast_es_on_breach_days"]).number_format == "0.00%"
    assert ws.cell(first, hdr["breach_rate"]).number_format == "0.00%"
    unformatted = []
    for sheet in wb.worksheets:
        if sheet.title == "Methodology":
            continue
        for row in sheet.iter_rows(min_row=5):
            for c in row:
                if isinstance(c.value, float) and c.number_format == "General":
                    unformatted.append((sheet.title, c.coordinate, c.value))
    assert not unformatted, unformatted[:5]


# ---- CLI ----
def test_parse_weights():
    w = cli.parse_weights("VTI=60,AGG=40")
    assert w["VTI"] == pytest.approx(0.6) and w["AGG"] == pytest.approx(0.4) and w["GLD"] == 0
    with pytest.raises(ValueError, match="Unknown ticker"):
        cli.parse_weights("ZZZ=100")
    with pytest.raises(ValueError, match="sum to"):
        cli.parse_weights("VTI=60,AGG=30")


def test_cli_export_and_refusals(tmp_path, capsys):
    out = tmp_path / "r.xlsx"
    assert cli.main(["export-excel", "--synthetic", "--out", str(out), "--weights", "VTI=50,TLT=50"]) == 0 and out.exists()
    assert cli.main(["brief", "--synthetic", "--out", str(tmp_path / "b.md")]) == 1
    assert "SYNTHETIC" in capsys.readouterr().err
    assert cli.main(["export-excel", "--csv", str(tmp_path / "nope.csv"), "--out", str(out)]) == 1
    assert cli.main(["export-excel", "--synthetic", "--weights", "VTI=10", "--out", str(out)]) == 1


def test_cli_download_failure_exit_code(monkeypatch, tmp_path, capsys):
    def offline(t, s, e):
        raise RuntimeError("no network")
    monkeypatch.setattr(cli, "download_prices", lambda *a, **k: ing.download_prices(*a, **{**k, "downloader": offline, "cache_dir": tmp_path}))
    monkeypatch.setattr(cli, "download_fred", lambda *a, **k: ing.download_fred(*a, fetcher=lambda *x: (_ for _ in ()).throw(OSError("x")), cache_dir=tmp_path))
    assert cli.main(["download", "--start", "2020-01-01"]) == 2
    assert "UNAVAILABLE" in capsys.readouterr().err


def test_cli_missing_weighted_ticker_is_reported(tmp_path, prices):
    csv = tmp_path / "p.csv"
    prices.drop(columns=["GLD"]).to_csv(csv)
    with pytest.raises(SystemExit, match="GLD"):
        cli.main(["export-excel", "--csv", str(csv), "--out", str(tmp_path / "o.xlsx")])
    assert cli.main(["export-excel", "--csv", str(csv), "--weights", "VTI=60,AGG=40", "--out", str(tmp_path / "o.xlsx")]) == 0


# ---- model-risk demonstration is reproducible and shows the intended result ----
def test_model_risk_demo_result():
    import model_risk_demo as d
    df = d.run().set_index("model")
    assert df.loc["parametric VaR", "Kupiec p"] < 0.001                     # normal VaR rejected on fat tails
    assert df.loc["parametric VaR", "breaches"] > 1.5 * df.loc["parametric VaR", "expected"]
    assert df.loc["historical VaR", "Kupiec p"] > 0.05
    assert df.loc["parametric VaR", "loss on breach / forecast ES"] > 1.0   # normal ES understates tail losses


def test_beta_ols_matches_covariance_beta_and_gives_interval():
    rng = np.random.default_rng(5)
    b = pd.Series(rng.normal(0, 0.01, 1000))
    a = 0.8 * b + pd.Series(rng.normal(0, 0.005, 1000))
    o = an.beta_ols(a, b)
    assert o["beta"] == pytest.approx(an.beta(a, b))
    assert o["beta_ci95_low"] < 0.8 < o["beta_ci95_high"] and 0 < o["r_squared"] < 1
    with pytest.raises(ValueError):
        an.beta_ols(a.iloc[:10], b.iloc[:10])
