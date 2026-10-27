"""Excel (.xlsx, openpyxl) export and research-brief rendering.

The workbook stores Python-computed VALUES. A few cells (weight totals, contribution
sums) are live formulas as cross-checks; they may differ from Python outputs only by
floating-point/display rounding. Compatible with LibreOffice Calc (no macros/tables).
"""
from __future__ import annotations

import datetime as dt
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from .pipeline import AnalysisResult

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(bold=True, color="FFFFFF")
WARN_FILL = PatternFill("solid", fgColor="FFF2CC")
DISCLAIMER = "Educational historical research. Not investment advice. Past results do not predict future results."

PCT, PCT1, NUM2, NUM4, DATE, INT = "0.00%", "0.0%", "0.00", "0.0000", "yyyy-mm-dd", "#,##0"

# Explicit column -> number format. Unlisted numeric columns stay "General" rather than
# being guessed from their names (a guessed format once displayed a 1.5% loss as "0").
_COLUMN_FORMATS: dict[str, str] = {
    **dict.fromkeys(["ann_return_arith", "ann_volatility", "ann_volatility_cov_check", "downside_deviation",
                     "max_drawdown", "worst_day", "best_day", "current_drawdown", "weight", "marginal_contribution",
                     "component_contribution", "var", "es", "breach_rate", "expected_breach_rate",
                     "breach_rate_ci95_low", "breach_rate_ci95_high", "mean_forecast_var",
                     "mean_realised_loss_on_breach", "mean_forecast_es_on_breach_days", "shock", "contribution",
                     "approx_price_return", "event_return", "contribution_buy_hold",
                     "portfolio_return_daily_rebalanced", "portfolio_return_buy_and_hold",
                     "max_drawdown_in_window", "worst_day_in_window"], PCT),
    **dict.fromkeys(["confidence", "pct_contribution"], PCT1),
    **dict.fromkeys(["sharpe", "sortino", "beta_vs_benchmark", "corr_vs_benchmark", "kupiec_lr",
                     "christoffersen_ind_lr", "expected_breaches", "years_of_history", "missing_pct_inside_span",
                     "duration_assumed", "convexity_assumed", "min_pct", "max_pct"], NUM2),
    **dict.fromkeys(["kupiec_p_value", "christoffersen_ind_p_value"], NUM4),
    **dict.fromkeys(["n_obs", "n_forecasts", "breaches", "tail_observations", "n_missing_inside_span",
                     "n_missing", "n_nonpositive", "n_large_return_flags", "trading_days",
                     "days_trough_to_recovery"], INT),
}


def _fmt_for(col: str, value) -> str | None:
    """Number format for a cell, or None (General) when the column is not registered."""
    if isinstance(value, (pd.Timestamp, dt.date, dt.datetime)):
        return DATE
    if isinstance(value, bool) or not isinstance(value, (int, float, np.integer, np.floating)):
        return None
    return _COLUMN_FORMATS.get(col)


def _clean_value(v):
    if isinstance(v, pd.Timestamp):
        return v.to_pydatetime().date()
    if isinstance(v, (np.floating, float)) and (np.isnan(v) or np.isinf(v)):
        return None
    if isinstance(v, np.integer):
        return int(v)
    if isinstance(v, np.floating):
        return float(v)
    if isinstance(v, np.bool_):
        return bool(v)
    if v is pd.NaT:
        return None
    return v


def _write_df(ws, df: pd.DataFrame, row: int, index_label: str | None = None) -> tuple[int, int]:
    """Write df with styled header at ``row``. Returns (first_data_row, last_data_row)."""
    d = df.reset_index() if index_label is not None else df
    if index_label is not None:
        d = d.rename(columns={d.columns[0]: index_label})
    for j, col in enumerate(d.columns, 1):
        c = ws.cell(row=row, column=j, value=str(col))
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(horizontal="center", wrap_text=True)
    for i, (_, rec) in enumerate(d.iterrows(), row + 1):
        for j, col in enumerate(d.columns, 1):
            v = _clean_value(rec[col])
            c = ws.cell(row=i, column=j, value=v)
            f = _fmt_for(str(col), v if not isinstance(v, dt.date) else pd.Timestamp(v))
            if f and v is not None:
                c.number_format = f
    return row + 1, row + len(d)


def _title(ws, title: str, note: str, source: str) -> None:
    ws["A1"], ws["A1"].font = title, Font(bold=True, size=14, color="1F3864")
    ws["A2"], ws["A2"].font = note, Font(italic=True, color="595959")
    ws["A3"], ws["A3"].font = source, Font(italic=True, color="595959")


def _autowidth(ws, minimum: int = 12, maximum: int = 60) -> None:
    widths: dict[int, int] = {}
    for row in ws.iter_rows(min_row=4):
        for c in row:
            if c.value is not None:
                widths[c.column] = max(widths.get(c.column, 0), len(str(c.value)) if not isinstance(c.value, float) else 10)
    for col, w in widths.items():
        ws.column_dimensions[get_column_letter(col)].width = max(minimum, min(maximum, w + 2))


def _stress_summary(res: AnalysisResult) -> pd.DataFrame:
    rows = []
    for r in res.stress_replays:
        rows.append({"scenario": r["event"], "type": "Historical replay (observed)", "start": r["start"], "end": r["end"],
                     "portfolio_return_daily_rebalanced": r.get("portfolio_return_daily_rebalanced"),
                     "portfolio_return_buy_and_hold": r.get("portfolio_return_buy_and_hold"),
                     "max_drawdown_in_window": r.get("max_drawdown_in_window"), "worst_day": r.get("worst_day"),
                     "trough_date": r.get("trough_date"), "recovery_date": r.get("recovery_date"),
                     "status": "available" if r["portfolio_available"] else f"UNAVAILABLE: {r['reason']}"})
    for h in res.stress_hypothetical:
        rows.append({"scenario": h["scenario"], "type": "HYPOTHETICAL shock (not history)",
                     "portfolio_return_daily_rebalanced": h["portfolio_return"], "status": "hypothetical"})
    return pd.DataFrame(rows)


METHODOLOGY_LINES = [
    "Returns: daily simple returns p_t/p_{t-1}-1 from adjusted closes; missing prices give missing returns (never zero).",
    "Portfolio: constant weights rebalanced daily; no transaction costs, taxes, fees or liquidity effects.",
    "Annualisation: arithmetic mean x 252; volatility x sqrt(252). Risk-free rate = annual input / 252 per day.",
    "Risk contribution: Euler decomposition of annualised portfolio volatility; components sum to portfolio volatility.",
    "VaR/ES: 1-day, positive loss thresholds. Historical = empirical quantile / tail mean; parametric = normal.",
    "Backtest: rolling window; each forecast uses only returns before the forecast target date. Breach = loss > VaR.",
    "Kupiec tests unconditional coverage only; low expected breach counts give low power. ES is not validated by a VaR test.",
    "Stress: historical replay uses observed prices only; hypothetical shocks and duration approximations are labelled as such.",
    "Treasury yields (FRED DGS10/DGS2) are context variables, not bond total returns.",
    "Data: yfinance is an unofficial source; ETF history is short; survivorship and look-back biases apply.",
    "Workbook formulas and Python outputs may differ by display/floating-point rounding.",
    DISCLAIMER,
]


def export_excel(res: AnalysisResult, path: str | Path) -> Path:
    """Write the formatted multi-sheet workbook. Returns the output path."""
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    meta = res.metadata
    src = (f"Source: {meta.get('source', 'n/a')} | Retrieved (UTC): {meta.get('retrieved_at_utc', 'n/a')}"
           + (" | SYNTHETIC DEMO DATA - NOT MARKET DATA" if meta.get("synthetic") else ""))
    wb = Workbook()
    ws = wb.active
    ws.title = "Executive Summary"
    _title(ws, "Multi-Asset Portfolio Risk Review", DISCLAIMER, src)
    p, i = res.portfolio_stats, res.sample_info
    summary = pd.DataFrame([
        ("Sample start", i["start"]), ("Sample end", i["end"]), ("Trading-day observations", i["rows_after"]),
        ("Rows dropped for missing returns", i["rows_dropped_missing"]),
        ("Annualised arithmetic return", p["ann_return_arith"]), ("Annualised volatility", p["ann_volatility"]),
        ("Sharpe ratio", p["sharpe"]), ("Sortino ratio", p["sortino"]), ("Max drawdown", p["max_drawdown"]),
        ("Beta vs benchmark", p.get("beta_vs_benchmark")),
        ("Beta 95% CI low (Newey-West)", p.get("beta_ci95_low")), ("Beta 95% CI high (Newey-West)", p.get("beta_ci95_high")),
        ("R-squared vs benchmark", p.get("r_squared")), ("Risk-free rate (annual)", res.assumptions["rf_annual"]),
        ("Data quality status", res.quality.get("status", "n/a")),
    ], columns=["Metric", "Value"])
    r0, r1 = _write_df(ws, summary, 5)
    fm = {"Annualised arithmetic return": PCT, "Annualised volatility": PCT, "Max drawdown": PCT,
          "Sharpe ratio": NUM2, "Sortino ratio": NUM2, "Beta vs benchmark": NUM2, "Beta 95% CI low (Newey-West)": NUM2,
          "Beta 95% CI high (Newey-West)": NUM2, "R-squared vs benchmark": NUM2, "Risk-free rate (annual)": PCT,
          "Trading-day observations": INT, "Rows dropped for missing returns": INT, "Sample start": DATE, "Sample end": DATE}
    for r in range(r0, r1 + 1):
        f = fm.get(ws.cell(r, 1).value)
        if f:
            ws.cell(r, 2).number_format = f
    ws.cell(r1 + 2, 1, "Units: returns, VaR, ES and drawdowns are decimal fractions shown as percentages; VaR/ES are positive losses.")
    _autowidth(ws)
    ws.column_dimensions["A"].width = 44

    ws = wb.create_sheet("Weights")
    _title(ws, "Portfolio weights (hypothetical, user-defined)", DISCLAIMER, src)
    wdf = res.weights.rename("weight").to_frame()
    r0, r1 = _write_df(ws, wdf, 5, "ticker")
    ws.cell(r1 + 1, 1, "Total (formula)").font = Font(bold=True)
    ws.cell(r1 + 1, 2, f"=SUM(B{r0}:B{r1})").number_format = PCT
    _autowidth(ws)

    ws = wb.create_sheet("Asset Statistics")
    _title(ws, "Asset-level and portfolio statistics (each on its own valid history)", DISCLAIMER, src)
    st = res.asset_stats.copy()
    st.loc["PORTFOLIO"] = pd.Series(res.portfolio_stats).reindex(st.columns)
    r0, r1 = _write_df(ws, st, 5, "series")
    _autowidth(ws)
    ws.freeze_panes = "B6"

    ws = wb.create_sheet("Correlation Matrix")
    _title(ws, f"Correlation of daily returns ({i['start']:%Y-%m-%d} to {i['end']:%Y-%m-%d}, complete-case dates)", DISCLAIMER, src)
    r0, r1 = _write_df(ws, res.corr.round(6), 5, "ticker")
    for r in range(r0, r1 + 1):
        for c in range(2, res.corr.shape[1] + 2):
            ws.cell(r, c).number_format = NUM2
    ws.conditional_formatting.add(f"B{r0}:{get_column_letter(res.corr.shape[1] + 1)}{r1}",
                                  ColorScaleRule(start_type="num", start_value=-1, start_color="F8696B",
                                                 mid_type="num", mid_value=0, mid_color="FFFFFF",
                                                 end_type="num", end_value=1, end_color="63BE7B"))
    _autowidth(ws)

    ws = wb.create_sheet("Risk Contributions")
    _title(ws, "Volatility risk contributions (Euler decomposition, annualised)", DISCLAIMER, src)
    r0, r1 = _write_df(ws, res.risk_contrib, 5, "ticker")
    ws.cell(r1 + 1, 1, "Total (formula)").font = Font(bold=True)
    for col in (2, 4, 5):
        L = get_column_letter(col)
        ws.cell(r1 + 1, col, f"=SUM({L}{r0}:{L}{r1})").number_format = PCT
    ws.cell(r1 + 3, 1, f"Portfolio volatility (annualised): {res.portfolio_stats['ann_volatility']:.4%}; "
                       "component contributions sum to this value.")
    _autowidth(ws)

    ws = wb.create_sheet("VaR and ES")
    _title(ws, "1-day VaR and Expected Shortfall (positive losses, full sample, in-sample)", DISCLAIMER, src)
    v = res.var_table.copy()
    _write_df(ws, v, 5)
    ws.cell(len(v) + 7, 1, "Historical ES = mean of losses at or beyond historical VaR. Parametric = normal distribution.")
    _autowidth(ws)

    ws = wb.create_sheet("Backtesting")
    _title(ws, f"Rolling out-of-sample VaR backtest (window = {res.assumptions['backtest_window']} days)", DISCLAIMER, src)
    if res.backtest_summary.empty:
        ws["A5"] = "Not run: insufficient observations for the chosen window."
    else:
        r0, r1 = _write_df(ws, res.backtest_summary, 5)
        hdr = {ws.cell(5, c).value: c for c in range(1, ws.max_column + 1)}
        if "kupiec_p_value" in hdr:
            L = get_column_letter(hdr["kupiec_p_value"])
            ws.conditional_formatting.add(f"{L}{r0}:{L}{r1}", CellIsRule(operator="lessThan", formula=["0.05"], fill=WARN_FILL))
        ws.cell(r1 + 2, 1, "Highlighted Kupiec p-value < 0.05 = coverage rejected at 5%. A pass does not prove reliability; "
                           "ES is descriptive only.")
    _autowidth(ws)
    ws.freeze_panes = "A6"

    ws = wb.create_sheet("Stress Scenarios")
    _title(ws, "Stress scenarios: historical replay vs hypothetical shocks (kept separate)", DISCLAIMER, src)
    r0, r1 = _write_df(ws, _stress_summary(res), 5)
    row = r1 + 3
    for rp in res.stress_replays:
        ws.cell(row, 1, f"Asset returns: {rp['event']}").font = Font(bold=True)
        if len(rp["asset_table"]):
            _, last = _write_df(ws, rp["asset_table"], row + 1, "asset")
            row = last + 2
        else:
            ws.cell(row + 1, 1, rp["reason"] or "n/a")
            row += 3
    for h in res.stress_hypothetical:
        ws.cell(row, 1, f"HYPOTHETICAL: {h['scenario']}").font = Font(bold=True, color="C00000")
        _, last = _write_df(ws, h["table"], row + 1, "asset")
        row = last + 2
    _autowidth(ws)

    ws = wb.create_sheet("Data Quality")
    _title(ws, "Data quality diagnostics", DISCLAIMER, src)
    q = res.quality
    if q:
        pt = pd.DataFrame(q.get("per_ticker", {})).T.drop(columns=["large_gaps", "large_return_flags"], errors="ignore")
        r0, r1 = _write_df(ws, pt, 5, "ticker")
        if q.get("yields"):
            ws.cell(r1 + 2, 1, "Treasury yields (FRED, context only)").font = Font(bold=True)
            _, r1 = _write_df(ws, pd.DataFrame(q["yields"]).T, r1 + 3, "series")
        iss = pd.DataFrame(q.get("issues", [])) if q.get("issues") else pd.DataFrame({"message": ["No issues detected"]})
        ws.cell(r1 + 2, 1, f"Status: {q.get('status')}").font = Font(bold=True)
        _write_df(ws, iss, r1 + 3)
    else:
        ws["A5"] = "No quality report attached."
    _autowidth(ws)

    ws = wb.create_sheet("Methodology")
    _title(ws, "Methodology and limitations", DISCLAIMER, src)
    for k, line in enumerate(METHODOLOGY_LINES, 5):
        ws.cell(k, 1, line).alignment = Alignment(wrap_text=True, vertical="top")
    ws.column_dimensions["A"].width = 130
    wb.save(path)
    return path


def render_brief(res: AnalysisResult, template_path: str | Path) -> str:
    """Populate the research-brief template with calculated numbers. Refuses synthetic data."""
    if res.metadata.get("synthetic"):
        raise ValueError("Refusing to populate the research brief from SYNTHETIC data.")
    if res.quality.get("status") == "FAIL":
        raise ValueError("Data quality status is FAIL; resolve data problems before writing the brief.")
    text = Path(template_path).read_text(encoding="utf-8")
    p, i = res.portfolio_stats, res.sample_info
    hist99 = res.var_table.query("model == 'historical' and confidence == 0.99")
    top = res.risk_contrib.sort_values("pct_contribution", ascending=False).iloc[0]
    vals = {
        "SAMPLE_START": f"{i['start']:%Y-%m-%d}", "SAMPLE_END": f"{i['end']:%Y-%m-%d}", "N_OBS": f"{i['rows_after']:,}",
        "ANN_RETURN": f"{p['ann_return_arith']:.2%}", "ANN_VOL": f"{p['ann_volatility']:.2%}",
        "MAX_DD": f"{p['max_drawdown']:.2%}", "SHARPE": f"{p['sharpe']:.2f}", "BETA": f"{p.get('beta_vs_benchmark', float('nan')):.2f}",
        "HIST_VAR99": f"{hist99['var'].iloc[0]:.2%}" if len(hist99) else "n/a",
        "HIST_ES99": f"{hist99['es'].iloc[0]:.2%}" if len(hist99) else "n/a",
        "TOP_RISK_ASSET": str(top.name), "TOP_RISK_PCT": f"{top['pct_contribution']:.1%}",
        "TOP_RISK_WEIGHT": f"{top['weight']:.1%}",
        "WEIGHTS": ", ".join(f"{k} {v:.0%}" for k, v in res.weights.items() if v > 0),
        "SOURCE": str(res.metadata.get("source", "n/a")), "RETRIEVED": str(res.metadata.get("retrieved_at_utc", "n/a")),
        "DQ_STATUS": str(res.quality.get("status", "n/a")),
    }
    bt = res.backtest_summary
    vals["BACKTEST_TABLE"] = ("\n".join(f"- {r.model} {r.confidence:.0%}: {r.breaches} breaches in {r.n_forecasts} forecasts "
                                        f"({r.breach_rate:.2%} vs {r.expected_breach_rate:.2%} expected; Kupiec p={r.kupiec_p_value:.3f})"
                                        for r in bt.itertuples()) if len(bt) else "Not run.")
    stress = []
    for r in res.stress_replays:
        stress.append(f"- {r['event']}: " + (f"{r['portfolio_return_daily_rebalanced']:.2%}" if r["portfolio_available"]
                                             else f"unavailable ({r['reason']})"))
    vals["STRESS_TABLE"] = "\n".join(stress)
    for k, v in vals.items():
        text = text.replace("{{" + k + "}}", v)
    return text
